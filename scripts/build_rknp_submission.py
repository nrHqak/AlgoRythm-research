#!/usr/bin/env python3
"""Build the final РКНП submission package.

Pipeline
--------
1. Read ``paper/rknp_src/RKNP_BODY_TEMPLATE_RU.md`` (authored text carrying
   ``[[bibkey]]`` citation tokens) and ``paper/rknp_src/references.json``.
2. Assign reference numbers by FIRST APPEARANCE, rewrite every token to ``[n]``,
   and emit the numbered reference list. Fails loudly on an orphan reference or
   an undefined key, so the two can never drift apart.
3. Write the fully resolved Markdown to ``paper/RKNP_SUBMISSION_FINAL_SOURCE.md``.
4. Render the same block list to
   ``paper/RKNP_SUBMISSION_FINAL_RU.pdf`` (ReportLab) and
   ``paper/RKNP_SUBMISSION_FINAL_RU.docx`` (python-docx) with identical page
   geometry, font, size and line spacing.
5. Report the rendered page span of every top-level section so the РКНП page
   budget can be verified against actual rendered pages, not estimates.

No scientific value is computed here; every number comes from the template,
which was written against the frozen analysis artifacts.
"""

from __future__ import annotations

import json
import pathlib
import re
import sys
from dataclasses import dataclass, field

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

ROOT = pathlib.Path(__file__).resolve().parents[1]
SRC = ROOT / "paper" / "rknp_src"
PAPER = ROOT / "paper"
FIGDIR = PAPER / "figures" / "rknp_final"
NIS_LOGO = SRC / "assets" / "nis_logo.png"

TEMPLATE = SRC / "RKNP_BODY_TEMPLATE_RU.md"
REFS_JSON = SRC / "references.json"
RESOLVED_MD = PAPER / "RKNP_SUBMISSION_FINAL_SOURCE.md"
OUT_PDF = PAPER / "RKNP_SUBMISSION_FINAL_RU.pdf"
OUT_DOCX = PAPER / "RKNP_SUBMISSION_FINAL_RU.docx"

# ---- page geometry (shared by both renderers) -----------------------------
PAGE_W, PAGE_H = A4
MARGIN = 2.0 * cm
BODY_PT = 12
LINE_SPACING = 1.15
LEADING = BODY_PT * LINE_SPACING
INDENT_CM = 1.25
TEXT_W = PAGE_W - 2 * MARGIN

SYSFONTS = pathlib.Path("/System/Library/Fonts/Supplemental")
FONT_FILES = {
    "TNR": "Times New Roman.ttf",
    "TNR-Bold": "Times New Roman Bold.ttf",
    "TNR-Italic": "Times New Roman Italic.ttf",
    "TNR-BoldItalic": "Times New Roman Bold Italic.ttf",
}
DOCX_FONT = "Times New Roman"


def register_fonts() -> None:
    for name, fname in FONT_FILES.items():
        path = SYSFONTS / fname
        if not path.exists():
            sys.exit(f"missing font: {path}")
        pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily(
        "TNR", normal="TNR", bold="TNR-Bold",
        italic="TNR-Italic", boldItalic="TNR-BoldItalic",
    )


# ===========================================================================
# 1-2. Citation resolution
# ===========================================================================
CITE_RE = re.compile(r"\[\[([A-Za-z0-9_|]+)\]\]")


def resolve_citations(text: str, entries: dict[str, str]) -> tuple[str, list[str]]:
    order: list[str] = []

    def repl(m: re.Match) -> str:
        keys = m.group(1).split("|")
        nums = []
        for k in keys:
            if k not in entries:
                sys.exit(f"citation key not in references.json: {k}")
            if k not in order:
                order.append(k)
            nums.append(order.index(k) + 1)
        return "[" + ", ".join(str(n) for n in nums) + "]"

    resolved = CITE_RE.sub(repl, text)

    unused = [k for k in entries if k not in order]
    if unused:
        sys.exit(f"orphan reference(s), cited nowhere: {unused}")
    if CITE_RE.search(resolved):
        sys.exit("unresolved citation token remains")

    ref_list = [f"{i + 1}. {entries[k]}" for i, k in enumerate(order)]
    return resolved, ref_list


# ===========================================================================
# 3. Block parser
# ===========================================================================
@dataclass
class Block:
    kind: str
    text: str = ""
    items: list = field(default_factory=list)
    meta: dict = field(default_factory=dict)


def parse_blocks(md: str, ref_list: list[str]) -> list[Block]:
    lines = md.split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped == "@@TITLEPAGE@@":
            tp: list[str] = []
            i += 1
            while lines[i].strip() != "@@ENDTITLEPAGE@@":
                tp.append(lines[i].rstrip())
                i += 1
            blocks.append(Block("titlepage", items=tp))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if re.fullmatch(r"@@TABLE[A-ZА-Я0-9]*@@", stripped):
            i += 1
            continue

        if stripped == "@@TOC@@":
            blocks.append(Block("toc"))
            i += 1
            continue

        if stripped == "@@REFERENCES@@":
            for r in ref_list:
                blocks.append(Block("ref", text=r))
            i += 1
            continue

        m = re.fullmatch(r"@@FIG\|([^|]+)\|([\d.]+)\|(.+)@@", stripped)
        if m:
            blocks.append(Block("fig", text=m.group(3),
                                meta={"file": m.group(1),
                                      "width_cm": float(m.group(2))}))
            i += 1
            continue

        if stripped.startswith("### "):
            blocks.append(Block("h2", text=stripped[4:].strip()))
            i += 1
            continue

        if stripped.startswith("## "):
            blocks.append(Block("h2", text=stripped[3:].strip()))
            i += 1
            continue

        if stripped.startswith("# "):
            blocks.append(Block("h1", text=stripped[2:].strip()))
            i += 1
            continue

        # table title line immediately followed by a pipe table
        if (stripped.startswith("**Таблица") and i + 1 < n
                and lines[i + 1].strip().startswith("|")):
            title = stripped
            i += 1
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", text=title, items=rows))
            continue

        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", text="", items=rows))
            continue

        if re.match(r"^- ", stripped):
            items = []
            while i < n and re.match(r"^- ", lines[i].strip()):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            blocks.append(Block("bullets", items=items))
            continue

        if re.match(r"^\d+\. ", stripped):
            items = []
            while i < n and re.match(r"^\d+\. ", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s*", "", lines[i].strip()))
                i += 1
            blocks.append(Block("numbered", items=items))
            continue

        kind = "note" if stripped.startswith("*Примечание.*") else "p"
        blocks.append(Block(kind, text=stripped))
        i += 1

    return blocks


# ===========================================================================
# Inline markup
# ===========================================================================
STAR = "\u0000"          # placeholder for an escaped literal asterisk


def inline_runs(text: str) -> list[tuple[str, bool, bool, bool]]:
    """Recursive inline parser -> [(text, bold, italic, mono)].

    Handles emphasis nested inside strong emphasis (``**величина *p***``),
    which a flat regex tokenizer mis-splits.
    """
    text = text.replace("\\*", STAR)

    def walk(t: str, bold: bool, italic: bool):
        runs: list[tuple[str, bool, bool, bool]] = []
        buf = ""
        i, n = 0, len(t)

        def flush():
            nonlocal buf
            if buf:
                runs.append((buf.replace(STAR, "*"), bold, italic, False))
                buf = ""

        while i < n:
            if t.startswith("**", i):
                j = t.find("**", i + 2)
                while j != -1 and j + 2 < n and t[j + 2] == "*":
                    j += 1          # "***x* y**": the closer is the last pair
                if j != -1 and j > i + 2:
                    flush()
                    runs.extend(walk(t[i + 2:j], True, italic))
                    i = j + 2
                    continue
            if t[i] == "*":
                j = t.find("*", i + 1)
                if j != -1 and j > i + 1:
                    flush()
                    runs.extend(walk(t[i + 1:j], bold, True))
                    i = j + 1
                    continue
            if t[i] == "`":
                j = t.find("`", i + 1)
                if j != -1 and j > i + 1:
                    flush()
                    runs.append((t[i + 1:j].replace(STAR, "*"),
                                 bold, italic, True))
                    i = j + 1
                    continue
            buf += t[i]
            i += 1
        flush()
        return runs

    return walk(text, False, False)



def esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def rl_markup(text: str) -> str:
    parts = []
    for t, b, it, mono in inline_runs(text):
        s = esc(t)
        if mono:
            s = f'<font face="Courier" size="10.5">{s}</font>'
        if it:
            s = f"<i>{s}</i>"
        if b:
            s = f"<b>{s}</b>"
        parts.append(s)
    return "".join(parts)


# ===========================================================================
# 4a. PDF renderer
# ===========================================================================
class Doc(BaseDocTemplate):
    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self.headings: list[tuple[int, str, int]] = []

    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph):
            sn = flowable.style.name
            if sn in ("H1", "H2"):
                txt = re.sub(r"<[^>]+>", "", flowable.getPlainText())
                self.headings.append((1 if sn == "H1" else 2, txt, self.page))


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = dict(fontName="TNR", fontSize=BODY_PT, leading=LEADING,
                textColor=colors.black, allowWidows=0, allowOrphans=0)
    return {
        "body": ParagraphStyle("body", alignment=TA_JUSTIFY,
                               firstLineIndent=INDENT_CM * cm,
                               spaceAfter=0, **base),
        "note": ParagraphStyle("note", alignment=TA_JUSTIFY, fontName="TNR",
                               fontSize=10.5, leading=12.5, spaceBefore=2,
                               spaceAfter=6, textColor=colors.black),
        "list": ParagraphStyle("list", alignment=TA_JUSTIFY,
                               leftIndent=0.7 * cm, firstLineIndent=0,
                               spaceAfter=2, bulletIndent=0.1 * cm, **base),
        "H1": ParagraphStyle("H1", alignment=TA_CENTER, fontName="TNR-Bold",
                             fontSize=14, leading=17, spaceBefore=0,
                             spaceAfter=12, textColor=colors.black,
                             keepWithNext=1),
        "H2": ParagraphStyle("H2", alignment=TA_LEFT, fontName="TNR-Bold",
                             fontSize=12, leading=15, spaceBefore=7,
                             spaceAfter=2, textColor=colors.black,
                             keepWithNext=1),
        "tabtitle": ParagraphStyle("tabtitle", alignment=TA_LEFT,
                                   fontName="TNR", fontSize=11, leading=13,
                                   spaceBefore=8, spaceAfter=3,
                                   textColor=colors.black, keepWithNext=1),
        "cap": ParagraphStyle("cap", alignment=TA_CENTER, fontName="TNR",
                              fontSize=10.5, leading=12.5, spaceBefore=5,
                              spaceAfter=9, textColor=colors.black),
        "cell": ParagraphStyle("cell", alignment=TA_LEFT, fontName="TNR",
                               fontSize=9.5, leading=11.2,
                               textColor=colors.black),
        "cellh": ParagraphStyle("cellh", alignment=TA_LEFT,
                                fontName="TNR-Bold", fontSize=9.5,
                                leading=11.2, textColor=colors.black),
        "ref": ParagraphStyle("ref", alignment=TA_JUSTIFY, fontName="TNR",
                              fontSize=11, leading=13.2, spaceAfter=4,
                              leftIndent=0.7 * cm, firstLineIndent=-0.7 * cm,
                              textColor=colors.black),
        "toc1": ParagraphStyle("toc1", alignment=TA_LEFT, fontName="TNR-Bold",
                               fontSize=11, leading=14.2, spaceBefore=3,
                               textColor=colors.black),
        "toc2": ParagraphStyle("toc2", alignment=TA_LEFT, fontName="TNR",
                               fontSize=11, leading=14.2, leftIndent=0.8 * cm,
                               textColor=colors.black),
        "tp": ParagraphStyle("tp", alignment=TA_CENTER, fontName="TNR",
                             fontSize=13, leading=18, textColor=colors.black),
        "tptitle": ParagraphStyle("tptitle", alignment=TA_CENTER,
                                  fontName="TNR-Bold", fontSize=15,
                                  leading=21, textColor=colors.black),
        "tpsub": ParagraphStyle("tpsub", alignment=TA_CENTER, fontName="TNR",
                                fontSize=12, leading=16,
                                textColor=colors.black),
        "tpright": ParagraphStyle("tpright", alignment=TA_RIGHT,
                                  fontName="TNR", fontSize=12, leading=17,
                                  textColor=colors.black),
        "tpschool": ParagraphStyle("tpschool", alignment=TA_CENTER,
                                   fontName="TNR-Bold", fontSize=12,
                                   leading=16, textColor=colors.black),
        "tplogo": ParagraphStyle("tplogo", alignment=TA_CENTER,
                                 fontName="TNR", fontSize=8.5,
                                 leading=10, textColor=colors.HexColor("#555555")),
        "tpdirection": ParagraphStyle("tpdirection", alignment=TA_CENTER,
                                      fontName="TNR", fontSize=12,
                                      leading=16, textColor=colors.black),
    }


def col_widths(rows: list[list[str]], total: float,
               font: str = "TNR-Bold", size: float = 9.5,
               pad: float = 8.0) -> list[float]:
    """Column widths that never break a word if the table can avoid it."""
    ncol = max(len(r) for r in rows)
    mins, prefs = [], []
    for c in range(ncol):
        longest_word = 0.0
        longest_cell = 0.0
        for r in rows:
            cell = re.sub(r"[*`]", "", r[c]) if c < len(r) else ""
            longest_cell = max(longest_cell,
                               pdfmetrics.stringWidth(cell, font, size))
            for w in cell.split():
                longest_word = max(longest_word,
                                   pdfmetrics.stringWidth(w, font, size))
        mins.append(longest_word + pad)
        prefs.append(max(longest_cell + pad, longest_word + pad))

    if sum(mins) >= total:                       # cannot avoid breaking
        k = total / sum(mins)
        return [m * k for m in mins]

    slack = total - sum(mins)
    extra = [max(p - m, 0.0) for p, m in zip(prefs, mins)]
    tot_extra = sum(extra)
    if tot_extra <= 0:
        return [m + slack / ncol for m in mins]
    k = min(1.0, slack / tot_extra)
    widths = [m + e * k for m, e in zip(mins, extra)]
    left = total - sum(widths)
    if left > 0:                                  # spread the remainder
        widths = [w + left * (e + 1) / (tot_extra + ncol)
                  for w, e in zip(widths, extra)]
    return widths


def build_table(rows: list[list[str]], st: dict, total: float) -> Table:
    widths = col_widths(rows, total)
    data = []
    for ri, row in enumerate(rows):
        style = st["cellh"] if ri == 0 else st["cell"]
        data.append([Paragraph(rl_markup(c), style) for c in row])
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.93, 0.93, 0.93)),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    return t


def titlepage_sections(items: list[str]) -> dict[str, list[str]]:
    """Split the semantic title-page source into layout regions."""
    marker_to_mode = {
        "@@NIS_HEADER@@": "header",
        "@@TITLE@@": "title",
        "@@SUBTITLE@@": "sub",
        "@@RIGHT@@": "right",
        "@@DIRECTION@@": "direction",
        "@@CITY@@": "city",
    }
    sections = {name: [] for name in marker_to_mode.values()}
    mode = "header"
    for raw in items:
        s = raw.strip()
        if s in marker_to_mode:
            mode = marker_to_mode[s]
        elif s and s != "@@SEP@@":
            sections[mode].append(s)
    return sections


def render_titlepage(items: list[str], st: dict) -> list:
    """Render a screenshot-matched NIS title page with stable geometry."""
    if not NIS_LOGO.exists():
        sys.exit(f"missing NIS logo: {NIS_LOGO}")
    sections = titlepage_sections(items)

    logo = Image(str(NIS_LOGO), width=3.2 * cm, height=1.78 * cm)
    logo.hAlign = "CENTER"
    logo_block = [logo, Spacer(1, 0.08 * cm), Paragraph(
        "Nazarbayev<br/>Intellectual<br/>Schools", st["tplogo"])]
    school = "<br/>".join(rl_markup(s) for s in sections["header"])
    header = Table(
        [[logo_block, Paragraph(school, st["tpschool"])]],
        colWidths=[4.0 * cm, TEXT_W - 4.0 * cm],
        hAlign="LEFT",
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))

    flow: list = [Spacer(1, 0.15 * cm), header, Spacer(1, 3.0 * cm)]
    for s in sections["title"]:
        flow.append(Paragraph(rl_markup(s), st["tptitle"]))
    flow.append(Spacer(1, 0.55 * cm))
    for s in sections["sub"]:
        flow.append(Paragraph(rl_markup(s), st["tpsub"]))

    flow.append(Spacer(1, 1.55 * cm))
    for s in sections["right"]:
        flow.append(Paragraph(rl_markup(s), st["tpright"]))

    flow.append(Spacer(1, 2.4 * cm))
    for s in sections["direction"]:
        flow.append(Paragraph(rl_markup(s), st["tpdirection"]))

    flow.append(Spacer(1, 1.2 * cm))
    for s in sections["city"]:
        flow.append(Paragraph(rl_markup(s), st["tpdirection"]))
    return flow


def render_pdf(blocks: list[Block], toc_entries, page_of_first_body: int) -> Doc:
    st = pdf_styles()
    doc = Doc(str(OUT_PDF), pagesize=A4,
              leftMargin=MARGIN, rightMargin=MARGIN,
              topMargin=MARGIN, bottomMargin=MARGIN,
              title="Специфичные для алгоритмического паттерна структурные "
                    "подсказки в автоматической локализации логических ошибок",
              author="AlgoRythm research")
    frame = Frame(MARGIN, MARGIN, TEXT_W, PAGE_H - 2 * MARGIN, id="body",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    def on_page(canvas, _doc):
        canvas.saveState()
        if canvas.getPageNumber() > 1:
            canvas.setFont("TNR", 11)
            canvas.setFillColor(colors.black)
            canvas.drawCentredString(PAGE_W / 2, MARGIN * 0.45,
                                     str(canvas.getPageNumber()))
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="main", frames=[frame],
                                       onPage=on_page)])

    flow: list = []
    first_h1 = True
    pending: list = []
    held_heading = None

    def flush():
        nonlocal pending
        if pending:
            flow.append(KeepTogether(pending))
            pending = []

    for bi, b in enumerate(blocks):
        nxt = blocks[bi + 1] if bi + 1 < len(blocks) else None
        if b.kind == "titlepage":
            flow.extend(render_titlepage(b.items, st))
            continue
        if b.kind == "h1":
            flush()
            if not first_h1:
                flow.append(PageBreak())
            else:
                flow.append(PageBreak())
                first_h1 = False
            flow.append(Paragraph(rl_markup(b.text), st["H1"]))
            continue
        if b.kind == "h2":
            flush()
            para = Paragraph(rl_markup(b.text), st["H2"])
            if nxt is not None and nxt.kind in ("table", "fig"):
                held_heading = para      # keep it with the block it introduces
            else:
                flow.append(para)
            continue
        if b.kind == "toc":
            dot_w = pdfmetrics.stringWidth(".", "TNR", 11)
            for lvl, txt, pg in toc_entries:
                style = st["toc1"] if lvl == 1 else st["toc2"]
                font = "TNR-Bold" if lvl == 1 else "TNR"
                avail = TEXT_W - style.leftIndent
                num = str(pg)
                w_txt = pdfmetrics.stringWidth(txt, font, 11)
                w_num = pdfmetrics.stringWidth(num, font, 11)
                gap = avail - w_txt - w_num - 8
                if gap < 4 * dot_w:      # title wraps: no leader, plain layout
                    flow.append(Paragraph(f"{esc(txt)}&nbsp;&nbsp;{num}", style))
                    continue
                dots = "." * int(gap / dot_w)
                flow.append(Paragraph(
                    f"{esc(txt)}&nbsp;{dots}&nbsp;{num}", style))
            continue
        if b.kind == "p":
            flush()
            flow.append(Paragraph(rl_markup(b.text), st["body"]))
            continue
        if b.kind == "note":
            para = Paragraph(rl_markup(b.text), st["note"])
            if pending:                 # keep a table note with its table
                pending.append(para)
                flush()
            else:
                flow.append(para)
            continue
        if b.kind == "ref":
            flush()
            flow.append(Paragraph(rl_markup(b.text), st["ref"]))
            continue
        if b.kind in ("bullets", "numbered"):
            flush()
            for k, it in enumerate(b.items, 1):
                bullet = "—" if b.kind == "bullets" else f"{k}."
                flow.append(Paragraph(rl_markup(it), st["list"],
                                      bulletText=bullet))
            continue
        if b.kind == "table":
            flush()
            grp = []
            if held_heading is not None:
                grp.append(held_heading)
                held_heading = None
            if b.text:
                grp.append(Paragraph(rl_markup(b.text), st["tabtitle"]))
            grp.append(build_table(b.items, st, TEXT_W))
            pending = grp
            continue
        if b.kind == "fig":
            flush()
            png = FIGDIR / f"{b.meta['file']}.png"
            if not png.exists():
                sys.exit(f"missing figure: {png}")
            from PIL import Image as PILImage
            with PILImage.open(png) as im:
                w_px, h_px = im.size
            w = b.meta["width_cm"] * cm
            h = w * h_px / w_px
            img = Image(str(png), width=w, height=h)
            img.hAlign = "CENTER"
            grp = []
            if held_heading is not None:
                grp.append(held_heading)
                held_heading = None
            grp += [Spacer(1, 0.25 * cm), img,
                    Paragraph(rl_markup(b.text), st["cap"])]
            flow.append(KeepTogether(grp))
            continue

    flush()
    doc.build(flow)
    return doc


# ===========================================================================
# 4b. DOCX renderer
# ===========================================================================
def set_cell_font(cell, size=9.5, bold=False):
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            r.font.name = DOCX_FONT
            r.font.size = Pt(size)
            r.font.bold = bold or r.font.bold
            r.font.color.rgb = RGBColor(0, 0, 0)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)


def remove_table_borders(table) -> None:
    """Remove every visible border from a layout-only DOCX table."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def add_runs(par, text, size=BODY_PT, base_bold=False, base_italic=False):
    for t, b, it, mono in inline_runs(text):
        r = par.add_run(t)
        r.font.name = "Courier New" if mono else DOCX_FONT
        r.font.size = Pt(size - 1.5 if mono else size)
        r.font.bold = b or base_bold
        r.font.italic = it or base_italic
        r.font.color.rgb = RGBColor(0, 0, 0)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    return par


def fld(par, instr, text=None):
    """Insert a simple Word field with an optional cached result."""
    r1 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    fs = OxmlElement("w:fldChar")
    fs.set(qn("w:fldCharType"), "separate")
    r3._r.append(fs)
    if text is not None:
        par.add_run(text)
    r5 = par.add_run()
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    r5._r.append(fe)
    for r in par.runs:
        r.font.name = DOCX_FONT
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 0, 0)


def docx_toc(doc, toc_entries):
    """A real Word TOC field whose cached result is the correct static TOC."""
    paras = []
    for lvl, txt, pg in toc_entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0.0 if lvl == 1 else 0.8)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.0
        pf.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)
        paras.append((p, lvl, txt, pg))

    first_p = paras[0][0]
    # field begin + instruction + separate, prepended to the first entry
    pre = []
    r1 = first_p.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    pre.append(r1._r)
    r2 = first_p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r2._r.append(it)
    pre.append(r2._r)
    r3 = first_p.add_run()
    fs = OxmlElement("w:fldChar")
    fs.set(qn("w:fldCharType"), "separate")
    r3._r.append(fs)
    pre.append(r3._r)
    # w:pPr must stay the first child of w:p, so insert the field runs after it
    pPr = first_p._p.find(qn("w:pPr"))
    base = list(first_p._p).index(pPr) + 1 if pPr is not None else 0
    for el in reversed(pre):
        first_p._p.remove(el)
        first_p._p.insert(base, el)

    for p, lvl, txt, pg in paras:
        run = p.add_run(txt)
        run.font.name = DOCX_FONT
        run.font.size = Pt(11)
        run.font.bold = lvl == 1
        run.font.color.rgb = RGBColor(0, 0, 0)
        tr = p.add_run("\t" + str(pg))
        tr.font.name = DOCX_FONT
        tr.font.size = Pt(11)
        tr.font.color.rgb = RGBColor(0, 0, 0)

    last_p = paras[-1][0]
    r5 = last_p.add_run()
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    r5._r.append(fe)


def render_docx(blocks: list[Block], toc_entries, out_path=None,
                break_before_first_h1: bool = True) -> None:
    out_path = out_path or OUT_DOCX
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, Cm(2.0))

    normal = doc.styles["Normal"]
    normal.font.name = DOCX_FONT
    normal.font.size = Pt(BODY_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(INDENT_CM)
    pf.widow_control = True

    for name, size, bold, align, before, after in (
        ("Heading 1", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 12),
        ("Heading 2", 12, True, WD_ALIGN_PARAGRAPH.LEFT, 7, 2),
    ):
        s = doc.styles[name]
        s.font.name = DOCX_FONT
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = False
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
        spf = s.paragraph_format
        spf.alignment = align
        spf.first_line_indent = Cm(0)
        spf.space_before = Pt(before)
        spf.space_after = Pt(after)
        spf.line_spacing = 1.15
        spf.keep_with_next = True
        spf.widow_control = True

    # page number in the footer; the title page carries none
    sec.different_first_page_header_footer = True
    first_footer = sec.first_page_footer.paragraphs[0]
    first_footer.text = ""
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.first_line_indent = Cm(0)
    fld(footer_p, " PAGE ", "1")

    first_h1 = True

    def plain(text, style_key="body"):
        p = doc.add_paragraph()
        pfm = p.paragraph_format
        if style_key == "note":
            pfm.first_line_indent = Cm(0)
            pfm.space_before = Pt(2)
            pfm.space_after = Pt(6)
            add_runs(p, text, size=10.5)
        elif style_key == "cap":
            pfm.first_line_indent = Cm(0)
            pfm.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pfm.space_before = Pt(4)
            pfm.space_after = Pt(9)
            add_runs(p, text, size=10.5)
        elif style_key == "tabtitle":
            pfm.first_line_indent = Cm(0)
            pfm.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pfm.space_before = Pt(8)
            pfm.space_after = Pt(3)
            pfm.keep_with_next = True
            add_runs(p, text, size=11)
        elif style_key == "ref":
            pfm.first_line_indent = Cm(-0.7)
            pfm.left_indent = Cm(0.7)
            pfm.space_after = Pt(4)
            add_runs(p, text, size=11)
        else:
            add_runs(p, text)
        return p

    for b in blocks:
        if b.kind == "titlepage":
            if not NIS_LOGO.exists():
                sys.exit(f"missing NIS logo: {NIS_LOGO}")
            sections = titlepage_sections(b.items)

            header = doc.add_table(rows=1, cols=2)
            header.alignment = WD_TABLE_ALIGNMENT.LEFT
            header.autofit = False
            remove_table_borders(header)
            left, right = header.rows[0].cells
            left.width, right.width = Cm(4.0), Cm(13.0)
            left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            lp = left.paragraphs[0]
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lp.paragraph_format.first_line_indent = Cm(0)
            lp.paragraph_format.space_after = Pt(0)
            lp.add_run().add_picture(str(NIS_LOGO), width=Cm(3.2))
            logo_text = left.add_paragraph()
            logo_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
            logo_text.paragraph_format.first_line_indent = Cm(0)
            logo_text.paragraph_format.space_after = Pt(0)
            add_runs(logo_text, "Nazarbayev\nIntellectual\nSchools", size=8.5)

            rp = right.paragraphs[0]
            rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rp.paragraph_format.first_line_indent = Cm(0)
            rp.paragraph_format.line_spacing = 1.25
            rp.paragraph_format.space_after = Pt(0)
            for idx, s in enumerate(sections["header"]):
                if idx:
                    rp.add_run().add_break()
                add_runs(rp, s, size=12, base_bold=True)

            def title_paragraph(text, align, size, bold=False,
                                line_spacing=1.25, after=0):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.alignment = align
                p.paragraph_format.line_spacing = line_spacing
                p.paragraph_format.space_after = Pt(after)
                add_runs(p, text, size=size, base_bold=bold)
                return p

            spacer = doc.add_paragraph()
            spacer.paragraph_format.first_line_indent = Cm(0)
            spacer.paragraph_format.space_after = Pt(67)

            for s in sections["title"]:
                title_paragraph(s, WD_ALIGN_PARAGRAPH.CENTER, 15, bold=True,
                                line_spacing=1.3)
            title_paragraph("", WD_ALIGN_PARAGRAPH.CENTER, 12, after=8)
            for s in sections["sub"]:
                title_paragraph(s, WD_ALIGN_PARAGRAPH.CENTER, 12)

            title_paragraph("", WD_ALIGN_PARAGRAPH.CENTER, 12, after=25)
            for s in sections["right"]:
                title_paragraph(s, WD_ALIGN_PARAGRAPH.RIGHT, 12,
                                line_spacing=1.3)

            title_paragraph("", WD_ALIGN_PARAGRAPH.CENTER, 12, after=40)
            for s in sections["direction"]:
                title_paragraph(s, WD_ALIGN_PARAGRAPH.CENTER, 12)

            title_paragraph("", WD_ALIGN_PARAGRAPH.CENTER, 12, after=12)
            for s in sections["city"]:
                title_paragraph(s, WD_ALIGN_PARAGRAPH.CENTER, 12)
            continue

        if b.kind == "h1":
            if not first_h1 or break_before_first_h1:
                doc.add_page_break()
            first_h1 = False
            p = doc.add_paragraph(style="Heading 1")
            add_runs(p, b.text, size=14, base_bold=True)
            continue

        if b.kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            add_runs(p, b.text, size=12, base_bold=True)
            continue

        if b.kind == "toc":
            docx_toc(doc, toc_entries)
            continue

        if b.kind == "p":
            plain(b.text)
            continue
        if b.kind == "note":
            plain(b.text, "note")
            continue
        if b.kind == "ref":
            plain(b.text, "ref")
            continue

        if b.kind in ("bullets", "numbered"):
            for k, it in enumerate(b.items, 1):
                p = doc.add_paragraph()
                pfm = p.paragraph_format
                pfm.left_indent = Cm(0.75)
                pfm.first_line_indent = Cm(-0.45)
                pfm.space_after = Pt(2)
                marker = "— " if b.kind == "bullets" else f"{k}. "
                r = p.add_run(marker)
                r.font.name = DOCX_FONT
                r.font.size = Pt(BODY_PT)
                r.font.color.rgb = RGBColor(0, 0, 0)
                add_runs(p, it)
            continue

        if b.kind == "table":
            if b.text:
                plain(b.text, "tabtitle")
            rows = b.items
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            widths = col_widths(rows, TEXT_W)
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.width = Cm(widths[ci] / cm)
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.first_line_indent = Cm(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = 1.0
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    if ci < len(row):
                        add_runs(cp, row[ci], size=9.5,
                                 base_bold=(ri == 0))
            # repeat the header row on every page and forbid row splitting
            hdr = t.rows[0]
            trPr = hdr._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
            for r in t.rows:
                trPr2 = r._tr.get_or_add_trPr()
                cs = OxmlElement("w:cantSplit")
                trPr2.append(cs)
            continue

        if b.kind == "fig":
            png = FIGDIR / f"{b.meta['file']}.png"
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(png), width=Cm(b.meta["width_cm"]))
            plain(b.text, "cap")
            continue

    doc.save(str(out_path))


# ===========================================================================
# main
# ===========================================================================
def main() -> None:
    register_fonts()

    entries = json.loads(REFS_JSON.read_text(encoding="utf-8"))["entries"]
    template = TEMPLATE.read_text(encoding="utf-8")
    resolved, ref_list = resolve_citations(template, entries)

    blocks = parse_blocks(resolved, ref_list)

    # --- pass 1: no TOC entries yet, only to learn heading page numbers -----
    probe_toc = [(lvl, "x", 0) for lvl in
                 [b for b in []]]  # placeholder, replaced below
    headings_names = [(1 if b.kind == "h1" else 2, b.text)
                      for b in blocks if b.kind in ("h1", "h2")
                      and b.text != "ОГЛАВЛЕНИЕ"]
    probe_toc = [(lvl, txt, 0) for lvl, txt in headings_names]

    doc1 = render_pdf(blocks, probe_toc, 1)
    pages1 = {t: p for _, t, p in doc1.headings}

    # --- pass 2: real page numbers ------------------------------------------
    toc_entries = [(lvl, txt, pages1.get(txt, 0)) for lvl, txt in headings_names]
    doc2 = render_pdf(blocks, toc_entries, 1)
    pages2 = {t: p for _, t, p in doc2.headings}

    # --- pass 3 if pagination shifted (TOC length change) -------------------
    if pages2 != pages1:
        toc_entries = [(lvl, txt, pages2.get(txt, 0)) for lvl, txt in headings_names]
        doc3 = render_pdf(blocks, toc_entries, 1)
        pages2 = {t: p for _, t, p in doc3.headings}
        toc_entries = [(lvl, txt, pages2.get(txt, 0)) for lvl, txt in headings_names]

    render_docx(blocks, toc_entries)

    # --- resolved source ----------------------------------------------------
    header = (
        "<!-- GENERATED FILE — do not edit by hand.\n"
        "     Source of truth: paper/rknp_src/RKNP_BODY_TEMPLATE_RU.md\n"
        "                      paper/rknp_src/references.json\n"
        "     Rebuild:         .venv/bin/python scripts/build_rknp_submission.py\n"
        "     Citations below are already resolved to sequential [n] numbers\n"
        "     assigned by first appearance; the reference list is generated\n"
        "     from the same pass, so it can contain no orphan and no gap. -->\n\n"
    )
    out = resolved.replace("@@REFERENCES@@", "\n\n".join(ref_list))
    RESOLVED_MD.write_text(header + out, encoding="utf-8")

    # --- report --------------------------------------------------------------
    def span(a: str, b: str) -> int:
        return pages2[b] - pages2[a]

    total_pages = doc2.page if hasattr(doc2, "page") else 0
    print(f"references cited : {len(ref_list)}")
    print(f"pdf pages        : {max(pages2.values())}+ (last heading page)")
    print("--- rendered page spans ---")
    print(f"1. ВВЕДЕНИЕ                 : starts p.{pages2['1. ВВЕДЕНИЕ']}, "
          f"{span('1. ВВЕДЕНИЕ', '2. ИССЛЕДОВАТЕЛЬСКАЯ ЧАСТЬ')} page(s)")
    print(f"2. ИССЛЕДОВАТЕЛЬСКАЯ ЧАСТЬ  : starts p.{pages2['2. ИССЛЕДОВАТЕЛЬСКАЯ ЧАСТЬ']}, "
          f"{span('2. ИССЛЕДОВАТЕЛЬСКАЯ ЧАСТЬ', '3. ЗАКЛЮЧЕНИЕ')} page(s)")
    print(f"3. ЗАКЛЮЧЕНИЕ               : starts p.{pages2['3. ЗАКЛЮЧЕНИЕ']}, "
          f"{span('3. ЗАКЛЮЧЕНИЕ', 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ')} page(s)")
    print(f"wrote {OUT_PDF.name}, {OUT_DOCX.name}, {RESOLVED_MD.name}")


if __name__ == "__main__":
    main()
